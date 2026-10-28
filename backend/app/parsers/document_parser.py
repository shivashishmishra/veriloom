import io
import re
from typing import List, Dict, Any, Optional

class ParsedParagraph:
    def __init__(self, index: int, text: str, page_number: Optional[int] = 1, metadata: Optional[Dict[str, Any]] = None):
        self.index = index
        self.text = text
        self.page_number = page_number
        self.metadata = metadata or {}

    def to_dict(self):
        return {
            "index": self.index,
            "text": self.text,
            "page_number": self.page_number,
            "metadata": self.metadata
        }

class DocumentParser:
    @staticmethod
    def parse_text(content: str) -> List[ParsedParagraph]:
        # Split by double newlines to form clean paragraphs
        raw_paras = [p.strip() for p in re.split(r'\n\s*\n', content) if p.strip()]
        if not raw_paras and content.strip():
            raw_paras = [content.strip()]
        return [ParsedParagraph(i, p, page_number=1) for i, p in enumerate(raw_paras)]

    @staticmethod
    def parse_markdown(content: str) -> List[ParsedParagraph]:
        paragraphs = []
        raw_paras = [p.strip() for p in re.split(r'\n\s*\n', content) if p.strip()]
        for i, p in enumerate(raw_paras):
            meta = {}
            # Check for image alt-text: ![alt text](url)
            img_matches = re.findall(r'!\[([^\]]*)\]\(([^)]+)\)', p)
            if img_matches:
                meta["images"] = [{"alt": m[0], "url": m[1]} for m in img_matches]
            # Check for hex colors like #1D3DD4
            hex_colors = re.findall(r'#(?:[0-9a-fA-F]{3}){1,2}\b', p)
            if hex_colors:
                meta["hex_colors"] = hex_colors
            paragraphs.append(ParsedParagraph(i, p, page_number=1, metadata=meta))
        return paragraphs

    @staticmethod
    def parse_docx(file_bytes: bytes) -> List[ParsedParagraph]:
        try:
            import docx
            doc = docx.Document(io.BytesIO(file_bytes))
            paragraphs = []
            idx = 0
            for p in doc.paragraphs:
                txt = p.text.strip()
                if txt:
                    paragraphs.append(ParsedParagraph(idx, txt, page_number=1, metadata={"style": p.style.name if p.style else None}))
                    idx += 1
            # Also extract tables
            for table in doc.tables:
                for row in table.rows:
                    row_txt = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_txt:
                        paragraphs.append(ParsedParagraph(idx, row_txt, page_number=1, metadata={"table": True}))
                        idx += 1
            return paragraphs
        except Exception as e:
            raise ValueError(f"DOCX extraction failed: {str(e)}")

    @staticmethod
    def parse_pdf(file_bytes: bytes) -> List[ParsedParagraph]:
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(file_bytes))
            paragraphs = []
            idx = 0
            for page_num, page in enumerate(reader.pages, start=1):
                page_text = page.extract_text() or ""
                raw_paras = [p.strip() for p in re.split(r'\n\s*\n', page_text) if p.strip()]
                for p in raw_paras:
                    paragraphs.append(ParsedParagraph(idx, p, page_number=page_num))
                    idx += 1
            if not paragraphs:
                raise ValueError("No extractable text found in PDF (may be a scanned image without embedded text).")
            return paragraphs
        except Exception as e:
            raise ValueError(f"PDF extraction failed: {str(e)}")

    @classmethod
    def parse_document(cls, filename: str, content_bytes: bytes) -> List[ParsedParagraph]:
        lower_name = filename.lower()
        if lower_name.endswith('.pdf'):
            return cls.parse_pdf(content_bytes)
        elif lower_name.endswith('.docx'):
            return cls.parse_docx(content_bytes)
        elif lower_name.endswith(('.md', '.markdown')):
            return cls.parse_markdown(content_bytes.decode('utf-8', errors='ignore'))
        elif lower_name.endswith(('.html', '.htm')):
            text = content_bytes.decode('utf-8', errors='ignore')
            # Extract alt texts from <img> tags
            alts = re.findall(r'<img[^>]*alt=["\']([^"\']*)["\']', text, re.IGNORECASE)
            # Normalize block closing tags to paragraph breaks
            normalized = re.sub(r'<\s*\/\s*(?:p|div|h[1-6]|li|tr)\s*>', '\n\n', text, flags=re.IGNORECASE)
            normalized = re.sub(r'<\s*br\s*\/?\s*>', '\n\n', normalized, flags=re.IGNORECASE)
            clean_text = re.sub(r'<[^>]+>', ' ', normalized)
            paras = cls.parse_text(clean_text)
            if alts and paras:
                paras[0].metadata["img_alts"] = alts
            return paras
        else:
            return cls.parse_text(content_bytes.decode('utf-8', errors='ignore'))
